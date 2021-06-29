import xlsxwriter
import datetime
import pymysql.cursors
import base64
import io
import pandas as pd
from BI_Export.InfoObject import InfoObject
import openpyxl

class ReportWriter:

    timestamp = ""
    filename = ""

    workbook = ""

    qrCodesByStunde = ""
    benefitsByAnzahl = ""
    benefitsByGeschaeft = ""
    benefitsInGeschaeft = ""

    def __init__(self):
        timestamp = datetime.datetime.now().strftime("%d_%m_%y_%H_%M")
        self.filename = "BI_Daily_Report_" + str(timestamp) + ".xlsx"
        self.output = None
        self.xlsx_data = None

    def get_filename(self):
        return self.filename

    def createDailyReport(self, path):

        #Zusammansetzung des Pfads muss u.U.für RasPI angepasst werden
        filepath = path + self.filename
        self.output = io.BytesIO()
        self.workbook = xlsxwriter.Workbook(self.output)

        self.writeRawDataSheet()
        self.writeReportSheet()

        self.workbook.close()
        self.xlsx_data = self.output
        return self.create_download_xlsx()

    def create_download_xlsx(self):
        # https://en.wikipedia.org/wiki/Data_URI_scheme
        self.xlsx_data.seek(0)
        media_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        data = base64.b64encode(self.xlsx_data.read()).decode("utf-8")
        href_data_downloadable = f'data:{media_type};base64,{data}'
        self.output.close()
        return data

    def getDBconnection(self):
        #Anpassungen nötig!
        # connection = pymysql.connect(host='localhost',    # change host-ip if needed
        #                              port=3310,           # change port if needed
        #                              user='dummy_insert',
        #                              password='1234',
        #                              db='RheinBerg_QRCode',
        #                              cursorclass=pymysql.cursors.DictCursor)

        connection = pymysql.connect(host='175.20.0.128',    # change host-ip if needed
                                     port=3306,           # change port if needed
                                     user='dummy_insert',
                                     password='1234',
                                     db='RheinBerg_QRCode',
                                     cursorclass=pymysql.cursors.DictCursor)
        return connection

    def readQrCodesByStunde(self, cursor):
        sql = "SELECT hour(create_time), count(*) FROM `RheinBerg_QRCode`.`BI_Light` group by hour(create_time);"
        cursor.execute(sql)
        return cursor.fetchall()

    def readBenefitsByAnzahl(self, cursor):
        sql = "SELECT Benefit, count(*) FROM `RheinBerg_QRCode`.`BI_Light` GROUP BY Benefit;"
        cursor.execute(sql)
        return cursor.fetchall()

    def readBenefitsByGeschaeft(self, cursor):
        sql = "SELECT Geschaeft, count(*) FROM `RheinBerg_QRCode`.`BI_Light` group by Geschaeft;"
        cursor.execute(sql)
        return cursor.fetchall()

    def readBenefitsInGeschaeft(self, cursor):
        sql = "SELECT Benefit, count(*) FROM `RheinBerg_QRCode`.`BI_Light` where Geschaeft = 'Best Döner' group by Benefit;"
        cursor.execute(sql)
        return cursor.fetchall()

    def writeData(self, iObj, worksheet):
        for i, x in enumerate(iObj.data):
            pos = iObj.posX + str(i + 1)
            worksheet.write_row(pos, [(iObj.data[i])[iObj.field1], (iObj.data[i])[iObj.field2]])

    def writeRawDataSheet(self):
        # Worksheet für BI-Rohdaten
        worksheet = self.workbook.add_worksheet('BI_Data')

        # InfoObjects erstellen
        connection = self.getDBconnection()
        with connection:
            with connection.cursor() as cursor:
                self.qrCodesByStunde = InfoObject(self.readQrCodesByStunde(cursor))
                self.benefitsByAnzahl = InfoObject(self.readBenefitsByAnzahl(cursor))
                self.benefitsByGeschaeft = InfoObject(self.readBenefitsByGeschaeft(cursor))
                self.benefitsInGeschaeft = InfoObject(self.readBenefitsInGeschaeft(cursor))

        self.qrCodesByStunde.posX = 'A'
        self.qrCodesByStunde.field1 = 'hour(create_time)'
        self.qrCodesByStunde.field2 = 'count(*)'

        self.benefitsByAnzahl.posX = 'D'
        self.benefitsByAnzahl.field1 = 'Benefit'
        self.benefitsByAnzahl.field2 = 'count(*)'

        self.benefitsByGeschaeft.posX = 'G'
        self.benefitsByGeschaeft.field1 = 'Geschaeft'
        self.benefitsByGeschaeft.field2 = 'count(*)'

        self.benefitsInGeschaeft.posX = 'J'
        self.benefitsInGeschaeft.field1 = 'Benefit'
        self.benefitsInGeschaeft.field2 = 'count(*)'

        self.writeData(self.qrCodesByStunde, worksheet)
        self.writeData(self.benefitsByAnzahl, worksheet)
        self.writeData(self.benefitsByGeschaeft, worksheet)
        self.writeData(self.benefitsInGeschaeft, worksheet)


    def writeReportSheet(self):
        worksheet = self.workbook.add_worksheet('Report')

        format_header = self.workbook.add_format({'bold': True, 'font_size': 20, 'border': 0, 'bg_color': '#FFFFFF'})
        format_bg = self.workbook.add_format({'border': 0, 'bg_color': '#FFFFFF'})

        # Write a conditional format over a range.
        worksheet.conditional_format('A1:AD90', {'type': 'no_errors',
                                                 'format': format_bg})

        worksheet.write('A1', 'BI Daily-Report vom ' + datetime.datetime.now().strftime("%d.%m.%y"), format_header)

        chart1 = self.workbook.add_chart({'type': 'column'})
        chart1.add_series({
            'categories': ['BI_Data', 0, 0, self.qrCodesByStunde.countRows(), 0],
            'values': ['BI_Data', 0, 1, self.qrCodesByStunde.countRows(), 1]
        })
        chart1.height = 300
        chart1.width = 300
        chart1.chart_name = "QR Codes by Stunde"
        worksheet.insert_chart('B5', chart1)

        chart2 = self.workbook.add_chart({'type': 'pie'})
        chart2.add_series({
            'categories': ['BI_Data', 0, 3, self.benefitsByAnzahl.countRows(), 3],
            'values': ['BI_Data', 0, 4, self.benefitsByAnzahl.countRows(), 4]
        })
        chart2.height = 300
        chart2.width = 300
        chart2.chart_name = "Benefits by Anzahl"
        worksheet.insert_chart('G5', chart2)

        chart3 = self.workbook.add_chart({'type': 'bar'})
        chart3.add_series({
            'categories': ['BI_Data', 0, 6, self.benefitsByGeschaeft.countRows(), 6],
            'values': ['BI_Data', 0, 7, self.benefitsByGeschaeft.countRows(), 7]
        })
        chart3.height = 300
        chart3.width = 300
        chart3.chart_name = "Benefits by Geschäft"
        worksheet.insert_chart('L5', chart3)

        chart4 = self.workbook.add_chart({'type': 'pie'})
        chart4.add_series({
            'categories': ['BI_Data', 0, 9, self.benefitsInGeschaeft.countRows(), 9],
            'values': ['BI_Data', 0, 10, self.benefitsInGeschaeft.countRows(), 10]
        })
        chart4.height = 300
        chart4.width = 300
        chart4.chart_name = "Benefits in Geschäft"
        worksheet.insert_chart('Q5', chart4)
